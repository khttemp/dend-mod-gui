import os
import traceback


import program.sub.textSetting as textSetting
import program.sub.errorLogClass as errorLogClass
from tkinter import messagebox as mb

errObj = errorLogClass.ErrorLogObj()


def execSaveStageData(stageDataFile, denFile, quietFlag, importDict):
    from program.sub.ssUnity.SSDecrypt.denDecrypt import DenDecrypt
    import program.sub.ssUnity.ssUnityProcess as ssUnityProcess

    decryptFile = DenDecrypt(denFile)
    if not decryptFile.open():
        decryptFile.printError()
        mb.showerror(title=textSetting.textList["error"], message=textSetting.textList["errorList"]["E14"])
        return -1
    if len(decryptFile.allList) == 1:
        dataName = decryptFile.allList[0][0]
        if dataName != "stagedata":
            errMsg = textSetting.textList["errorList"]["E105"] + denFile
            mb.showerror(title=textSetting.textList["error"], message=errMsg)
            return -5
        data = decryptFile.allList[0][-1]
        if os.path.splitext(stageDataFile)[1].lower() != ".xlsx":
            script = ssUnityProcess.getScriptData(stageDataFile)
        else:
            rootPath = importDict["rootPath"]
            configPath = importDict["configPath"]
            result, obj = ssUnityProcess.loadExcelData(stageDataFile, data, rootPath, configPath)
            if not result:
                mb.showerror(title=textSetting.textList["error"], message=obj["message"])
                return -6
            script = ssUnityProcess.getScriptDataByExcel(obj["data"])
        ssUnityProcess.saveDenFile(data, decryptFile, script)
        if not quietFlag:
            mb.showinfo(title=textSetting.textList["success"], message=textSetting.textList["infoList"]["I51"])
        return 0
    else:
        errMsg = textSetting.textList["errorList"]["E105"] + denFile
        mb.showerror(title=textSetting.textList["error"], message=errMsg)
        return -6


def execSaveRail(excelFile, railFile, quietFlag, importDict):
    if os.path.splitext(railFile)[1].lower() == ".den":
        return execSaveStageData(excelFile, railFile, quietFlag, importDict)

    import openpyxl
    import program.sub.railEditor.dendDecrypt.RSdecrypt as dendRs
    import program.sub.railEditor.dendDecrypt.CSdecrypt as dendCs
    import program.sub.railEditor.dendDecrypt.BSdecrypt as dendBs
    import program.sub.railEditor.dendDecrypt.LSdecrypt as dendLs
    import program.sub.railEditor.dendDecrypt.RSExcelWidget as RSExcelWidget
    import program.sub.railEditor.dendDecrypt.CSExcelWidget as CSExcelWidget
    import program.sub.railEditor.dendDecrypt.BSExcelWidget as BSExcelWidget
    import program.sub.railEditor.dendDecrypt.LSExcelWidget as LSExcelWidget

    try:
        wb = openpyxl.load_workbook(excelFile, data_only=True)
        tabList = textSetting.textList["railEditor"]["railComboValue"]
        # ver
        ws = wb[tabList[0]]
        ver = ws.cell(1, 1).value

        oldVersionList = [
            "DEND_MAP_VER0100", "DEND_MAP_VER0101",
            "DEND_MAP_VER0102",
            "DEND_MAP_VER0110",
            "DEND_MAP_VER0300", "DEND_MAP_VER0400"
        ]

        if ver in oldVersionList:
            if ver == "DEND_MAP_VER0300" or ver == "DEND_MAP_VER0400":
                decryptFile = dendRs.RailDecrypt(railFile)
                excelWidget = RSExcelWidget.ExcelWidget(excelFile, decryptFile, importDict["configPath"])
            elif ver == "DEND_MAP_VER0110":
                decryptFile = dendCs.RailDecrypt(railFile)
                excelWidget = CSExcelWidget.ExcelWidget(excelFile, decryptFile, importDict["configPath"])
            elif ver == "DEND_MAP_VER0102":
                decryptFile = dendBs.RailDecrypt(railFile)
                excelWidget = BSExcelWidget.ExcelWidget(excelFile, decryptFile, importDict["configPath"])
            elif ver == "DEND_MAP_VER0100" or ver == "DEND_MAP_VER0101":
                decryptFile = dendLs.RailDecrypt(railFile)
                excelWidget = LSExcelWidget.ExcelWidget(excelFile, decryptFile, importDict["configPath"])

            if decryptFile.game == "LS":
                tabList = textSetting.textList["railEditor"]["railLsComboValue"]

            for tabName in tabList:
                if tabName not in wb.sheetnames:
                    errMsg = textSetting.textList["errorList"]["E95"].format(tabName)
                    mb.showerror(title=textSetting.textList["error"], message=errMsg)
                    return -3

            result, obj = excelWidget.loadExcelData()
            if not result:
                mb.showerror(title=textSetting.textList["error"], message=obj["message"])
                return -4
            excelWidget.saveRailFile(decryptFile.filePath, obj["data"])
            if not quietFlag:
                mb.showinfo(title=textSetting.textList["success"], message=textSetting.textList["infoList"]["I114"])
            return 0
    except Exception:
        errObj.write(traceback.format_exc())
        mb.showerror(title=textSetting.textList["error"], message=textSetting.textList["errorList"]["E14"])
        return -1


def cuiMain(argv, importDict):
    if argv[1] in ["/saveRail", "/quietSaveRail", "/debugSaveRail"]:
        quietFlag = False
        debugFlag = False
        if argv[1] == "/quietSaveRail":
            quietFlag = True
        elif argv[1] == "/debugSaveRail":
            debugFlag = True
        excelFile = argv[2]
        if not os.path.exists(excelFile):
            errMsg = textSetting.textList["errorList"]["E103"] + excelFile
            mb.showerror(title=textSetting.textList["error"], message=errMsg)
            return -2
        railFile = argv[3]
        if os.path.splitext(railFile)[1].lower() == ".bin":
            if not os.path.exists(railFile):
                w = open(railFile, "wb")
                w.close()
        else:
            if not os.path.exists(railFile):
                errMsg = textSetting.textList["errorList"]["E104"] + railFile
                mb.showerror(title=textSetting.textList["error"], message=errMsg)
                return -2
        if debugFlag:
            debugStr = "エクセルのパス\n{0}".format(os.path.abspath(excelFile))
            debugStr += "\n"
            debugStr += "ファイルのパス\n{0}".format(os.path.abspath(railFile))
            debugStr += "\nで確認しました。このまま実行しますか？"
            result = mb.askyesno(title=textSetting.textList["confirm"], message=debugStr)
            if not result:
                return 0
        return execSaveRail(excelFile, railFile, quietFlag, importDict)
